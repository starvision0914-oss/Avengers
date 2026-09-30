"""팔린 키워드 → 지마켓 신규광고센터 '직접 운영형' 키워드 등록용 엑셀 생성(읽기 전용 산출물, 2026-09-30).

구광고센터 키워드 리포트(GmarketKeywordReport)에 쌓인 '구매가 발생한 키워드'를 골라, 신규광고센터의
엑셀 대량등록 양식(sample_modification_광고그룹추가등록_직접운영형 = 기존 직접운영 캠페인에 그룹+키워드 추가)에
맞춰 계정별 파일로 만든다. 광고센터에 접속하거나 업로드하지 않으며 DB도 바꾸지 않는다 — 파일만 생성하고
업로드는 사람이 한다(먼저 *_TEST.xlsx 소량으로 확인 후 본 파일 업로드 권장).

선정 규칙(기본값):
  · 키워드 단위 구매 1건 이상, 광고수익률(ROAS) 450% 이상(지마켓 손익분기 추정 450~490%)
  · 지금 '판매중'인 상품만(판매중지/판매불가 제외), 도매마트 L코드 품절·미확인 상품 제외
  · 겹치는 기간 버킷(연간/롤링/월별)은 (계정,상품,키워드)별로 가장 큰 행 하나만 사용(이중계산 방지)
  · 입찰가 = 과거 평균클릭비용 × 1.1 (10원 단위 올림), 최소 100원 ~ 최대 1,000원
    ⚠ 실제 최저/최고 입찰가와 그룹당 상품 수 제한은 확인하지 못했다 — TEST 파일로 먼저 검증할 것.

  python manage.py gmarket_sold_keyword_excel
  python manage.py gmarket_sold_keyword_excel --accounts rejoice666 --min-roas 300
"""
import csv
import datetime as dt
import math
import os
import shutil

from django.core.management.base import BaseCommand
from django.db.models import Sum

OUT_ROOT = '/home/rejoice888/PUBLIC/지마켓_신규광고센터_분석'
TEMPLATE = os.path.join(OUT_ROOT, '엑셀등록_템플릿', 'sample_modification_광고그룹추가등록_직접운영형.xlsx')
SHEET_GROUP = '광고그룹추가등록_직접운영형'
SHEET_KW = '키워드별입찰가추가등록_직접운영형'
SHEET_PROD = '상품별입찰가추가등록_직접운영형'
PLACEHOLDER_CAMPAIGN = '★기존 직접운영 캠페인명을 입력하세요'


def _ceil10(v):
    return int(math.ceil(v / 10.0) * 10)


class Command(BaseCommand):
    help = '팔린 키워드를 신규광고센터 직접운영형 엑셀 등록 양식으로 계정별 생성(파일만, 읽기 전용)'

    def add_arguments(self, parser):
        parser.add_argument('--accounts', nargs='*')
        parser.add_argument('--min-roas', type=float, default=450.0)
        parser.add_argument('--min-orders', type=int, default=1)
        parser.add_argument('--bid-mult', type=float, default=1.1)
        parser.add_argument('--min-bid', type=int, default=100)
        parser.add_argument('--max-bid', type=int, default=1000)
        parser.add_argument('--group-size', type=int, default=10, help='광고그룹 1개에 넣을 상품 수')
        parser.add_argument('--area', default='검색결과', help='노출영역(검색결과 / 검색결과+추천영역)')

    def handle(self, *args, **o):
        import openpyxl
        from apps.cpc.models import (CrawlerAccount, GmarketKeywordReport as R, GmarketMyProduct as G,
                                     GmarketNewAdCost as C)
        from apps.cpc.eleven_my_product_service import get_lcode_soldout_rows

        if not os.path.exists(TEMPLATE):
            self.stderr.write(f'템플릿 없음: {TEMPLATE}')
            return
        today = dt.date.today()
        out_dir = os.path.join(OUT_ROOT, f'팔린키워드_등록엑셀_{today:%Y-%m-%d}')
        os.makedirs(out_dir, exist_ok=True)

        masters = [a.login_id for a in CrawlerAccount.objects.filter(platform='gmarket', is_active=True)
                   .order_by('display_order', 'login_id')
                   if not (a.gmarket_origin_id and a.gmarket_origin_id != a.login_id)]
        if o['accounts']:
            masters = [m for m in masters if m in o['accounts']]

        on_sale = set(G.objects.filter(market='gmarket', status_type='판매중').values_list('product_no', flat=True))
        lsold = {pno for _, pno in get_lcode_soldout_rows(G, 'seller_product_code', 'product_no', 'status_type', '판매중')}

        # 계정별 기존 직접 운영형 캠페인명(신규광고센터 캠페인 목록에서 수집된 값)
        camp = {}
        for lid, name in (C.objects.filter(campaign_type__contains='직접').values_list('login_id', 'campaign_name')
                          .distinct().order_by('login_id', 'campaign_name')):
            camp.setdefault(lid, []).append(name)

        summary = []
        for lid in masters:
            best = {}
            for r in R.objects.filter(login_id=lid, orders__gte=o['min_orders']).values(
                    'product_no', 'keyword', 'orders', 'conv_amount', 'cost', 'avg_click_cost'):
                kw = (r['keyword'] or '').strip()
                if not kw or r['cost'] <= 0:
                    continue
                key = (r['product_no'], kw)
                if key not in best or r['conv_amount'] > best[key]['conv_amount']:
                    best[key] = r
            rows = []
            skipped = {'roas': 0, 'not_on_sale': 0, 'lcode': 0}
            for (pno, kw), r in best.items():
                roas = r['conv_amount'] * 100.0 / r['cost']
                if roas < o['min_roas']:
                    skipped['roas'] += 1
                    continue
                if pno not in on_sale:
                    skipped['not_on_sale'] += 1
                    continue
                if pno in lsold:
                    skipped['lcode'] += 1
                    continue
                base = r['avg_click_cost'] or 0
                bid = min(o['max_bid'], max(o['min_bid'], _ceil10(base * o['bid_mult']))) if base else o['min_bid']
                rows.append((pno, kw, bid, r['orders'], r['conv_amount'], round(roas)))
            if not rows:
                summary.append([lid, 0, 0, 0, '', skipped['roas'], skipped['not_on_sale'], skipped['lcode']])
                continue
            rows.sort(key=lambda x: (-x[4], x[0], x[1]))
            pnos = []
            for r in rows:
                if r[0] not in pnos:
                    pnos.append(r[0])
            gname = {}
            for i, p in enumerate(pnos):
                gname[p] = f'팔린키워드_{today:%y%m%d}_{i // o["group_size"] + 1:02d}'
            groups = sorted(set(gname.values()))
            campaign = (next((n for n in camp.get(lid, []) if '판매된상품' in n), None)
                        or (camp.get(lid) or [PLACEHOLDER_CAMPAIGN])[0])

            def build(path, kw_rows, group_names):
                shutil.copy(TEMPLATE, path)
                wb = openpyxl.load_workbook(path)
                for name in (SHEET_GROUP, SHEET_KW, SHEET_PROD):
                    ws = wb[name]
                    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
                        for c in row:
                            c.value = None
                for i, g in enumerate(group_names, 3):
                    ws = wb[SHEET_GROUP]
                    ws.cell(i, 1, lid); ws.cell(i, 2, campaign); ws.cell(i, 3, g); ws.cell(i, 4, o['area'])
                for i, (pno, kw, bid, *_rest) in enumerate(kw_rows, 3):
                    ws = wb[SHEET_KW]
                    ws.cell(i, 1, lid); ws.cell(i, 2, gname[pno]); ws.cell(i, 3, pno); ws.cell(i, 4, kw); ws.cell(i, 5, bid)
                wb.save(path)

            build(os.path.join(out_dir, f'{lid}_직접운영_키워드등록.xlsx'), rows, groups)
            first_g = groups[0]
            test_rows = [r for r in rows if gname[r[0]] == first_g][:5]
            test_g = sorted({gname[r[0]] for r in test_rows})
            build(os.path.join(out_dir, f'{lid}_TEST_소량검증용.xlsx'), test_rows, test_g)
            summary.append([lid, len(rows), len(pnos), len(groups), campaign, skipped['roas'], skipped['not_on_sale'], skipped['lcode']])

        with open(os.path.join(out_dir, '_요약.csv'), 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(['계정', '키워드행', '상품수', '광고그룹수', '기존 직접운영 캠페인명', '제외(ROAS미달)', '제외(판매중아님)', '제외(L코드품절)'])
            w.writerows(summary)
        tot_rows = sum(s[1] for s in summary)
        need_camp = [s[0] for s in summary if s[1] and s[4] == PLACEHOLDER_CAMPAIGN]
        self.stdout.write(self.style.SUCCESS(
            f'저장: {out_dir}\n계정 {sum(1 for s in summary if s[1])}개 · 키워드 {tot_rows:,}행 · 캠페인명 직접 입력 필요 계정 {len(need_camp)}개 {need_camp}'))
